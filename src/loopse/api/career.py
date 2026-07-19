"""Career Navigator API - 职业导航、技能 GAP 分析、简历生成与模拟面试接口。"""
from __future__ import annotations

import json
import logging
import re
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, HTTPException
from pydantic import BaseModel, Field

from ..agent.career_navigator import career_navigator
from ..agent.profiler import ProfilerAgent

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/v1/career", tags=["career"])
_profiler = ProfilerAgent()


class GapAnalysisRequest(BaseModel):
    user_id: str = Field(default="student-001")
    target_job_id: str = Field(..., description="目标岗位 ID，如 job_001")


class CareerPathRequest(BaseModel):
    user_id: str = Field(default="student-001")
    target_job_id: str = Field(..., description="目标岗位 ID")


class ResumeGenerateRequest(BaseModel):
    user_id: str = Field(default="student-001")
    target_job: str = Field(default="网络工程师", description="目标岗位名称")
    extra_info: Optional[str] = Field(default=None, description="额外个人信息（项目经历、证书等）")


class ResumeOptimizeRequest(BaseModel):
    user_id: str = Field(default="student-001")
    resume_text: str = Field(..., min_length=50, description="现有简历文本")
    target_job: str = Field(default="网络工程师")
    job_description: Optional[str] = Field(default=None, description="目标岗位JD（可选）")


class MockInterviewRequest(BaseModel):
    user_id: str = Field(default="student-001")
    job_title: str = Field(default="网络工程师", description="岗位名称")
    job_type: Optional[str] = Field(default=None, description="岗位类型（兼容字段）")
    question_count: int = Field(default=8, ge=3, le=15)
    difficulty: str = Field(default="medium", description="easy/medium/hard")

    @property
    def effective_job_title(self) -> str:
        return self.job_type or self.job_title


class InterviewReportRequest(BaseModel):
    user_id: str = Field(default="student-001")
    job_title: str = Field(default="网络工程师")
    questions: list[dict] = Field(..., description="题目列表")
    answers: list[str] = Field(..., description="用户回答列表")


@router.get("/jobs")
def list_jobs():
    """获取所有支持的目标岗位列表。"""
    return {"jobs": career_navigator.list_jobs()}


@router.post("/gap-analysis")
def analyze_gap(req: GapAnalysisRequest):
    """分析用户与目标岗位之间的技能 GAP。"""
    try:
        profile = _profiler.get_profile(req.user_id)
        gap_report = career_navigator.analyze_gap(profile, req.target_job_id)
        if "error" in gap_report:
            raise HTTPException(status_code=400, detail=gap_report["error"])
        return gap_report
    except HTTPException:
        raise
    except Exception as exc:
        logger.error("[Career] GAP分析失败 user=%s job=%s: %s", req.user_id, req.target_job_id, exc)
        raise HTTPException(status_code=500, detail="GAP分析失败") from exc


@router.post("/path")
def get_career_path(req: CareerPathRequest):
    """根据目标岗位生成职业导向学习路径。"""
    try:
        profile = _profiler.get_profile(req.user_id)
        gap_report = career_navigator.analyze_gap(profile, req.target_job_id)
        if "error" in gap_report:
            raise HTTPException(status_code=400, detail=gap_report["error"])
        path = career_navigator.recommend_learning_path(
            user_id=req.user_id,
            user_profile=profile,
            target_job_id=req.target_job_id,
            gap_report=gap_report,
        )
        return {
            "gap_summary": {
                "job_name": gap_report["job_name"],
                "gap_score": gap_report["gap_score"],
                "gap_level": gap_report["gap_level"],
                "estimated_weeks": gap_report["estimated_weeks"],
            },
            "path": path,
        }
    except HTTPException:
        raise
    except Exception as exc:
        logger.error("[Career] 路径生成失败 user=%s job=%s: %s", req.user_id, req.target_job_id, exc)
        raise HTTPException(status_code=500, detail="职业路径生成失败") from exc


@router.post("/generate-resume")
def generate_resume(req: ResumeGenerateRequest):
    """根据用户画像自动生成 Markdown 格式简历。"""
    from ..core.llm_client import get_client_for_persona

    profile = _profiler.get_profile(req.user_id)
    dim_labels = {
        "conceptual_understanding": "概念理解", "protocol_analysis": "协议分析",
        "calculation_ability": "计算能力", "error_diagnosis": "错误诊断",
        "system_design": "系统设计", "knowledge_connection": "知识关联",
    }
    skills_text = "; ".join(
        f"{label}({round(profile.get(dim, 0.5) * 100)}%)"
        for dim, label in dim_labels.items()
        if profile.get(dim, 0.5) >= 0.6
    ) or "网络基础知识"
    strong_points = "、".join(profile.get("strong_points", [])[:5]) or "网络协议"
    extra = f"\n附加信息：{req.extra_info}" if req.extra_info else ""

    prompt = (
        f"请生成一份针对「{req.target_job}」岗位的计算机网络方向专业简历（Markdown格式）。\n\n"
        f"用户能力概况：\n"
        f"- 技能亮点：{skills_text}\n"
        f"- 掌握较好的知识点：{strong_points}\n"
        f"- 学习轮次：{profile.get('turn_count', 0)}轮{extra}\n\n"
        "简历应包含：\n"
        "1. 个人信息（占位符与说明）\n"
        "2. 求职意向（一句话，简洁有力）\n"
        "3. 技能清单（路由交换、协议栈、工具、认证）\n"
        "4. 项目经历（2-3个虚拟项目）\n"
        "5. 教育背景\n"
        "6. 自我评价\n"
        "以可直接修改的 Markdown 格式输出，标识【占位符】表示需填写内容。"
    )
    try:
        client = get_client_for_persona("expert")
        resume_md = client.chat(prompt, max_tokens=2000)
    except Exception as exc:
        logger.error("[Career] 简历生成失败: %s", exc)
        raise HTTPException(status_code=500, detail="简历生成失败")

    return {
        "user_id": req.user_id,
        "target_job": req.target_job,
        "generated_at": datetime.now().isoformat(),
        "resume_markdown": resume_md,
    }


@router.post("/optimize-resume")
def optimize_resume(req: ResumeOptimizeRequest):
    """一键优化简历文本，标注新增网络技能。"""
    from ..core.llm_client import get_client_for_persona

    profile = _profiler.get_profile(req.user_id)
    strong_points = "、".join(profile.get("strong_points", [])[:5]) or "网络协议"
    jd_text = f"\n岗位 JD：{req.job_description[:500]}" if req.job_description else ""

    prompt = (
        f"请优化以下针对「{req.target_job}」岗位的简历，突出计算机网络相关技能。{jd_text}\n\n"
        f"当前简历：\n{req.resume_text[:2000]}\n\n"
        f"用户实际擅长点：{strong_points}\n\n"
        "优化要求：\n"
        "1. 【添加】推断的网络技能条目（标注 NEW）\n"
        "2. 【添加】与岗位匹配的量化表述\n"
        "3. 【调整】语言更动词化、专业化\n"
        "4. 保留原有内容框架，只增强不削弱\n"
        "返回 Markdown 格式优化后简历。"
    )
    try:
        client = get_client_for_persona("expert")
        optimized = client.chat(prompt, max_tokens=2500)
    except Exception as exc:
        raise HTTPException(status_code=500, detail="简历优化失败")

    return {
        "user_id": req.user_id,
        "target_job": req.target_job,
        "optimized_resume_markdown": optimized,
    }


@router.post("/mock-interview")
def create_mock_interview(req: MockInterviewRequest):
    """生成模拟面试题组（按岗位类型）。"""
    from ..core.llm_client import get_client_for_persona

    profile = _profiler.get_profile(req.user_id)
    weak_points = "、".join(profile.get("weak_points", [])[:3]) or "网络协议层次"
    diff_map = {"easy": "入门", "medium": "中等", "hard": "进阶"}
    diff_label = diff_map.get(req.difficulty, "中等")
    job = req.effective_job_title

    prompt = (
        f"请生成{req.question_count}道针对『{job}』岗位的{diff_label}难度面试题。用户薄弱点：{weak_points}\n\n"
        f"题型要求：6道技术原理题 + 2道场景分析题（不足{req.question_count}道则全面覆盖网络栈知识）。\n"
        "返回 JSON 数组，每题包含：id(唯一字符串)、question(题目)、type(technical/scenario)、\n"
        "expected_key_points(核心考察点列表最多3个)、tips(面试小贴士)。只返回JSON。"
    )
    try:
        client = get_client_for_persona("expert")
        raw = client.chat(prompt, max_tokens=3000)
        clean = re.sub(r"```(?:json)?\s*|\s*```", "", raw).strip()
        questions = json.loads(clean)
        if not isinstance(questions, list):
            questions = []
    except Exception as exc:
        logger.error("[Career] 面试题生成失败: %s", exc)
        questions = [
            {"id": f"q_{i}", "question": f"请介绍一个{req.job_title}常见网络技术场景。",
             "type": "technical", "expected_key_points": ["网络协议"], "tips": "结合实际经验回答"}
            for i in range(min(req.question_count, 5))
        ]

    return {
        "user_id": req.user_id,
        "job_title": req.job_title,
        "difficulty": req.difficulty,
        "generated_at": datetime.now().isoformat(),
        "total_questions": len(questions),
        "questions": questions[:req.question_count],
    }


@router.post("/interview-report")
def generate_interview_report(req: InterviewReportRequest):
    """评分面试回答，生成三维度面试报告。"""
    from ..core.llm_client import get_client_for_persona

    if len(req.questions) != len(req.answers):
        raise HTTPException(status_code=400, detail="题目数量与回答数量不匹配")

    qa_pairs = ""
    for i, (q, a) in enumerate(zip(req.questions, req.answers), 1):
        qa_pairs += f"Q{i}: {q.get('question', '')[:200]}\nA{i}: {a[:200]}\n\n"

    prompt = (
        f"请对以下『{req.job_title}』面试的回答进行专业评分并生成报告。\n\n"
        f"面试问答：\n{qa_pairs}\n"
        "请以 JSON 格式返回评测结果：\n"
        '{"total_score": <0-100综合分>,'
        '"professional_score": <专业知识维度 0-100>,'
        '"expression_score": <语言表达维度 0-100>,'
        '"logic_score": <逻辑思维维度 0-100>,'
        '"strengths": ["亮点1", "亮点2"],'
        '"improvements": ["建议1", "建议2"],'
        '"overall_comment": "整体评语（不超过200字）",'
        '"recommended_resources": ["建议进修资源"]}'
        "\n只返回JSON。"
    )
    try:
        client = get_client_for_persona("expert")
        raw = client.chat(prompt, max_tokens=800)
        clean = re.sub(r"```(?:json)?\s*|\s*```", "", raw).strip()
        report_data = json.loads(clean)
    except Exception as exc:
        logger.error("[Career] 面试报告生成失败: %s", exc)
        report_data = {
            "total_score": 70, "professional_score": 70,
            "expression_score": 70, "logic_score": 70,
            "strengths": ["已提交面试答案"],
            "improvements": ["进一步完善技术表达"],
            "overall_comment": "已完成面试答题，建议继续加深网络协议学习。",
            "recommended_resources": ["计算机网络谢希仁教材"]
        }

    return {
        "user_id": req.user_id,
        "job_title": req.job_title,
        "question_count": len(req.questions),
        "reported_at": datetime.now().isoformat(),
        "report": report_data,
    }


# ── 前端适配别名路由 ──────────────────────────────────────────

class ResumeGenerateV2(BaseModel):
    user_id: str = "anonymous"
    name: str = ""
    target_job: str = ""
    skills: str = ""
    experience: str = ""
    profile_data: dict = {}


class ResumeEditRequest(BaseModel):
    current_resume: str
    instruction: str
    user_id: str = "anonymous"


class ResumeExportRequest(BaseModel):
    resume_text: str
    format: str = "docx"
    name: str = "简历"


@router.post("/resume/generate")
def resume_generate_v2(req: ResumeGenerateV2):
    """AI 一键生成简历，返回 preview 字段。"""
    from ..core.llm_client import get_client_for_persona

    profile = _profiler.get_profile(req.user_id) if req.user_id != "anonymous" else {}
    dim_labels = {
        "conceptual_understanding": "概念理解", "protocol_analysis": "协议分析",
        "system_design": "系统设计", "error_diagnosis": "错误诊断",
    }
    skills_from_profile = "; ".join(
        f"{label}({round(profile.get(dim, 0.5) * 100)}%)"
        for dim, label in dim_labels.items()
        if profile.get(dim, 0.5) >= 0.6
    ) or ""
    all_skills = "; ".join(filter(None, [req.skills, skills_from_profile])) or "计算机网络基础知识"

    prompt = (
        f"请生成一份专业的计算机网络方向简历（纯文本格式，方便直接显示）。\n\n"
        f"求职者信息：\n"
        f"- 姓名：{req.name or '【请填写】'}\n"
        f"- 目标岗位：{req.target_job or '网络工程师'}\n"
        f"- 技能：{all_skills}\n"
        f"- 实习/项目经历：{req.experience or '（未填写）'}\n\n"
        "简历包含：个人信息 | 求职意向 | 技能清单 | 项目经历(2-3个) | 教育背景 | 自我评价\n"
        "以纯文本格式输出，每个模块用《=====》分隔，占位内容用【...】标注。"
    )
    try:
        client = get_client_for_persona("expert")
        preview = client.chat(prompt, max_tokens=1500)
    except Exception as exc:
        logger.error("[Career] 简历生成v2失败: %s", exc)
        raise HTTPException(status_code=500, detail="简历生成失败")

    return {"preview": preview, "download_url": ""}


@router.post("/resume/edit")
def resume_edit(req: ResumeEditRequest):
    """自然语言指令修改简历内容。"""
    from ..core.llm_client import get_client_for_persona

    prompt = (
        f"请根据以下指令修改简历：\n\n"
        f"修改指令：{req.instruction}\n\n"
        f"当前简历：\n{req.current_resume[:3000]}\n\n"
        "请保留原有格式和结构，只按指令修改对应内容，返回修改后的完整简历文本。"
    )
    try:
        client = get_client_for_persona("expert")
        modified = client.chat(prompt, max_tokens=2000)
    except Exception as exc:
        raise HTTPException(status_code=500, detail="简历修改失败")

    return {"resume": modified}


@router.post("/resume/export")
def resume_export(req: ResumeExportRequest):
    """将简历文本导出为 docx 文件。"""
    import io
    from fastapi.responses import StreamingResponse

    if req.format == "docx":
        try:
            from docx import Document
            doc = Document()
            doc.add_heading(f"{req.name} 的简历", 0)
            for para in req.resume_text.split("\n"):
                if para.strip().startswith("====="):
                    doc.add_paragraph("─" * 40)
                elif para.strip():
                    doc.add_paragraph(para)
            buf = io.BytesIO()
            doc.save(buf)
            buf.seek(0)
            return StreamingResponse(
                buf,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": f'attachment; filename="resume_{req.name}.docx"'},
            )
        except ImportError:
            raise HTTPException(status_code=500, detail="python-docx 未安装")

    raise HTTPException(status_code=400, detail="PDF 导出需安装 reportlab，请先使用 Word 导出")
