"""内容导出 API — 将题目/画像/学习记录导出为 Word/Excel/PDF。"""
from __future__ import annotations

import io
import json
import logging
from typing import List, Optional

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/v1/export", tags=["export"])


# ── 数据模型 ────────────────────────────────────────────────────────────────

class QuestionItem(BaseModel):
    id: Optional[str] = None
    type: str = "choice"
    knowledge_point: str = ""
    question: str
    options: Optional[dict] = None
    answer: Optional[str] = None
    explanation: Optional[str] = None
    key_points: Optional[List[str]] = None
    sample_answer: Optional[str] = None
    raw_content: Optional[str] = None


class ExportQuestionsRequest(BaseModel):
    title: str = Field(default="计算机网络练习题")
    knowledge_point: str = Field(default="")
    questions: List[QuestionItem]
    format: str = Field(default="docx", pattern="^(docx|xlsx)$")
    include_answers: bool = Field(default=True)


class ExportProfileRequest(BaseModel):
    user_id: str
    username: str = "学生用户"
    format: str = Field(default="docx", pattern="^(docx|xlsx)$")


# ── Word 导出工具 ────────────────────────────────────────────────────────────

def _export_questions_docx(req: ExportQuestionsRequest) -> bytes:
    """将题目列表导出为 Word 文档（.docx）。"""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise HTTPException(status_code=500, detail="python-docx 未安装，请运行 pip install python-docx")

    doc = Document()

    # 标题
    title_para = doc.add_heading(req.title, 0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if req.knowledge_point:
        kp_para = doc.add_paragraph(f"知识点：{req.knowledge_point}")
        kp_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")  # 空行

    type_labels = {
        "choice": "选择题", "fill": "填空题", "short_answer": "简答题",
        "diagram": "图形判断", "case": "案例分析",
    }

    # 按题型分组
    groups: dict[str, list] = {}
    for q in req.questions:
        groups.setdefault(q.type, []).append(q)

    q_num = 1
    for q_type, items in groups.items():
        type_label = type_labels.get(q_type, q_type)
        doc.add_heading(f"一、{type_label}", level=1)

        for q in items:
            # 题干
            p = doc.add_paragraph()
            run = p.add_run(f"{q_num}. {q.question}")
            run.bold = True
            q_num += 1

            # 选项（选择题）
            if q.options and isinstance(q.options, dict):
                for opt_key, opt_val in q.options.items():
                    doc.add_paragraph(f"   {opt_key}. {opt_val}", style="List Bullet")

            # 填空/简答的提示空间
            if q.type in ("fill", "short_answer", "case"):
                doc.add_paragraph("")
                doc.add_paragraph("答：＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿＿")
                doc.add_paragraph("")

            # 答案与解析
            if req.include_answers:
                if q.answer:
                    ans_p = doc.add_paragraph()
                    ans_run = ans_p.add_run(f"【答案】{q.answer}")
                    ans_run.font.color.rgb = RGBColor(0x1a, 0x73, 0xe8)

                if q.explanation:
                    exp_p = doc.add_paragraph()
                    exp_run = exp_p.add_run(f"【解析】{q.explanation}")
                    exp_run.font.color.rgb = RGBColor(0x34, 0xa8, 0x53)

                if q.key_points:
                    kp_p = doc.add_paragraph()
                    kp_run = kp_p.add_run(f"【评分要点】" + " | ".join(q.key_points))
                    kp_run.font.color.rgb = RGBColor(0xfb, 0xbc, 0x04)

                if q.sample_answer:
                    sa_p = doc.add_paragraph()
                    sa_run = sa_p.add_run(f"【参考答案】{q.sample_answer}")
                    sa_run.font.color.rgb = RGBColor(0x34, 0xa8, 0x53)

            doc.add_paragraph("")  # 题目间距

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _export_questions_xlsx(req: ExportQuestionsRequest) -> bytes:
    """将题目列表导出为 Excel（.xlsx）。"""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl 未安装，请运行 pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "练习题"

    # 表头
    headers = ["序号", "题型", "知识点", "题目", "选项A", "选项B", "选项C", "选项D", "答案", "解析"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1a73e8")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    type_labels = {
        "choice": "选择题", "fill": "填空题", "short_answer": "简答题",
        "diagram": "图形判断", "case": "案例分析",
    }

    for idx, q in enumerate(req.questions, 1):
        opts = q.options or {}
        row = [
            idx,
            type_labels.get(q.type, q.type),
            q.knowledge_point,
            q.question,
            opts.get("A", ""),
            opts.get("B", ""),
            opts.get("C", ""),
            opts.get("D", ""),
            q.answer or ("|".join(q.key_points) if q.key_points else ""),
            q.explanation or q.sample_answer or "",
        ]
        ws.append(row)

    # 自适应列宽
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 60)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── 路由 ──────────────────────────────────────────────────────────────────

@router.post("/questions")
async def export_questions(req: ExportQuestionsRequest):
    """将题目列表导出为 Word 或 Excel 文件。"""
    if req.format == "docx":
        content = _export_questions_docx(req)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        filename = f"练习题_{req.knowledge_point or 'export'}.docx"
    else:
        content = _export_questions_xlsx(req)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"练习题_{req.knowledge_point or 'export'}.xlsx"

    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{filename}",
        "Content-Length": str(len(content)),
    }
    return StreamingResponse(io.BytesIO(content), media_type=media_type, headers=headers)


@router.post("/profile")
async def export_profile(req: ExportProfileRequest):
    """将用户学习档案导出为 Word 文档。"""
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise HTTPException(status_code=500, detail="python-docx 未安装")

    from ..agent.profiler import ProfilerAgent
    profiler = ProfilerAgent()
    profile = profiler.get_profile(req.user_id)

    doc = Document()
    title = doc.add_heading(f"{req.username} — 学习能力档案", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"用户ID：{req.user_id}    学习轮次：{profile.get('turn_count', 0)} 轮")
    doc.add_paragraph("")

    doc.add_heading("能力雷达（8维）", level=1)
    dim_labels = {
        "conceptual_understanding": "概念理解",
        "protocol_analysis": "协议分析",
        "calculation_ability": "计算能力",
        "error_diagnosis": "错误诊断",
        "system_design": "系统设计",
        "knowledge_connection": "知识关联",
        "expression_clarity": "表达清晰",
        "self_correction": "自我纠错",
    }
    for key, label in dim_labels.items():
        score = round(profile.get(key, 0.5) * 100)
        bar = "█" * (score // 10) + "░" * (10 - score // 10)
        doc.add_paragraph(f"{label}：{bar} {score}分")

    doc.add_paragraph("")
    doc.add_heading("薄弱知识点", level=1)
    for kp in profile.get("weak_points", []):
        doc.add_paragraph(f"• {kp}", style="List Bullet")

    doc.add_heading("已掌握知识点", level=1)
    for kp in profile.get("strong_points", []):
        doc.add_paragraph(f"• {kp}", style="List Bullet")

    buf = io.BytesIO()
    doc.save(buf)
    content = buf.getvalue()

    filename = f"学习档案_{req.username}.docx"
    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return StreamingResponse(
        io.BytesIO(content),
        media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


# ─── Path Export ──────────────────────────────────────────────────────────────

class PathNodeItem(BaseModel):
    node_id: str
    node_name: str
    chapter: Optional[str] = None
    status: Optional[str] = "pending"
    current_mastery: Optional[float] = 0.0
    recommendation_reason: Optional[str] = None


class ExportPathRequest(BaseModel):
    user_id: str
    username: str = "学生用户"
    path_title: str = "个性化学习路径"
    nodes: List[PathNodeItem]
    format: str = Field(default="docx", pattern="^(docx|xlsx)$")
    estimated_weeks: int = 0
    job_name: Optional[str] = None


def _export_path_docx(req: ExportPathRequest) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise HTTPException(status_code=500, detail="python-docx 未安装")

    doc = Document()
    title = doc.add_heading(req.path_title, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph(
        f"学习者：{req.username}   目标岗位：{req.job_name or '未指定'}   预计完成：{req.estimated_weeks} 周"
    )
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")

    doc.add_heading("学习路径节点", level=1)
    status_labels = {"completed": "✅ 已完成", "in_progress": "🔄 进行中", "pending": "⬜ 待学习"}
    for i, node in enumerate(req.nodes, 1):
        p = doc.add_paragraph()
        p.add_run(f"{i}. {node.node_name}").bold = True
        status_text = status_labels.get(node.status or "pending", node.status or "")
        mastery_pct = round((node.current_mastery or 0) * 100)
        doc.add_paragraph(f"   {status_text}  |  掌握度：{mastery_pct}%  |  章节：{node.chapter or '-'}")
        if node.recommendation_reason:
            doc.add_paragraph(f"   说明：{node.recommendation_reason}")
        doc.add_paragraph("")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _export_path_xlsx(req: ExportPathRequest) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl 未安装")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "学习路径"
    headers = ["序号", "节点名称", "章节", "状态", "当前掌握度", "说明"]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="10b981")
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    status_labels = {"completed": "已完成", "in_progress": "进行中", "pending": "待学习"}
    for idx, node in enumerate(req.nodes, 1):
        ws.append([
            idx,
            node.node_name,
            node.chapter or "",
            status_labels.get(node.status or "pending", node.status or ""),
            f"{round((node.current_mastery or 0) * 100)}%",
            node.recommendation_reason or "",
        ])
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.post("/path")
async def export_path(req: ExportPathRequest):
    """导出学习路径为 Word 或 Excel。"""
    if req.format == "docx":
        content = _export_path_docx(req)
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        filename = f"学习路径_{req.username}.docx"
    else:
        content = _export_path_xlsx(req)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"学习路径_{req.username}.xlsx"
    return StreamingResponse(
        io.BytesIO(content),
        media_type=media_type,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


# ─── Markdown Export ──────────────────────────────────────────────────────────

class ExportMarkdownRequest(BaseModel):
    title: str = "导出文档"
    content: str = Field(..., min_length=1, description="Markdown 内容")
    format: str = Field(default="docx", pattern="^(docx|pdf)$")


def _md_to_docx(title: str, md_content: str) -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        raise HTTPException(status_code=500, detail="python-docx 未安装")
    import re
    doc = Document()
    h = doc.add_heading(title, 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in md_content.splitlines():
        stripped = line.strip()
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        elif re.match(r"^\d+\. ", stripped):
            doc.add_paragraph(re.sub(r"^\d+\. ", "", stripped), style="List Number")
        elif stripped.startswith("```"):
            pass  # skip code fences
        elif stripped == "":
            doc.add_paragraph("")
        else:
            p = doc.add_paragraph()
            # handle **bold** inline
            parts = re.split(r"(\*\*[^*]+\*\*)", stripped)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@router.post("/markdown")
async def export_markdown(req: ExportMarkdownRequest):
    """将 Markdown 文本导出为 Word（.docx）或 PDF。"""
    if req.format == "pdf":
        # 尝试 weasyprint，降级到 Word
        try:
            from weasyprint import HTML
            import markdown
            html_body = markdown.markdown(req.content, extensions=["tables", "fenced_code"])
            full_html = f"""<!DOCTYPE html><html><head><meta charset='utf-8'>
            <style>body{{font-family:Arial,sans-serif;padding:2cm;font-size:12pt;}}
            h1,h2,h3{{color:#1e293b;}} pre{{background:#f1f5f9;padding:12px;border-radius:6px;}}
            code{{background:#f1f5f9;padding:2px 5px;border-radius:3px;}}
            table{{border-collapse:collapse;width:100%;}} th,td{{border:1px solid #ccc;padding:6px;}}
            </style></head><body><h1>{req.title}</h1>{html_body}</body></html>"""
            pdf_bytes = HTML(string=full_html).write_pdf()
            return StreamingResponse(
                io.BytesIO(pdf_bytes),
                media_type="application/pdf",
                headers={"Content-Disposition": f"attachment; filename*=UTF-8''{req.title}.pdf"},
            )
        except ImportError:
            # 降级到 docx
            content = _md_to_docx(req.title, req.content)
            return StreamingResponse(
                io.BytesIO(content),
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": f"attachment; filename*=UTF-8''{req.title}.docx"},
            )
    else:
        content = _md_to_docx(req.title, req.content)
        return StreamingResponse(
            io.BytesIO(content),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{req.title}.docx"},
        )
